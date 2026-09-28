"""Generate human-only DOCX and PDF annotations for a completed image benchmark."""
from __future__ import annotations

import json
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


PROJECT_ROOT = Path(__file__).resolve().parent
DATASET_PATH = PROJECT_ROOT / "dataset" / "vocabulary.json"
MANIFEST_PATH = PROJECT_ROOT / "results" / "manifest.jsonl"
REPORTS_PATH = PROJECT_ROOT / "reports"
DOCX_PATH = REPORTS_PATH / "ngaam_nou_image_benchmark_annotation.docx"
PDF_PATH = REPORTS_PATH / "ngaam_nou_image_benchmark_annotation.pdf"
PDF_CJK_FONT = "NgaamNouCJK"
PDF_CJK_FONT_PATH = Path(r"C:\Windows\Fonts\kaiu.ttf")

MODEL_ORDER = ("qwen/qwen-image-2", "google/nano-banana-2")
MODEL_LABELS = {
    "qwen/qwen-image-2": "Qwen Image 2",
    "google/nano-banana-2": "Nano Banana 2",
}
CRITERIA = (
    "Semantic clarity",
    "Child comprehension",
    "Target-word fidelity",
    "Visual simplicity",
    "Composition",
    "Text contamination",
    "Safety / age appropriateness",
    "Educational usefulness",
)


@dataclass(frozen=True)
class BenchmarkRun:
    """A complete, validated benchmark run selected from the manifest."""

    run_id: str
    visual_entries: dict[str, dict[str, dict[str, Any]]]
    grammar_entries: list[dict[str, Any]]


def load_dataset() -> list[dict[str, Any]]:
    """Load and validate the controlled vocabulary dataset."""
    records = json.loads(DATASET_PATH.read_text(encoding="utf-8"))
    if not isinstance(records, list) or len(records) != 10:
        raise ValueError("dataset/vocabulary.json must contain exactly 10 records")
    if not all(isinstance(record, dict) for record in records):
        raise ValueError("every dataset record must be a JSON object")
    return records


def load_manifest() -> list[dict[str, Any]]:
    """Load nonempty JSONL manifest records without modifying the source file."""
    if not MANIFEST_PATH.is_file():
        raise FileNotFoundError(f"benchmark manifest not found: {MANIFEST_PATH}")
    records = []
    for line_number, line in enumerate(MANIFEST_PATH.read_text(encoding="utf-8").splitlines(), 1):
        if not line.strip():
            continue
        try:
            record = json.loads(line)
        except json.JSONDecodeError as error:
            raise ValueError(f"invalid manifest JSON at line {line_number}") from error
        if not isinstance(record, dict):
            raise ValueError(f"manifest line {line_number} is not an object")
        records.append(record)
    if not records:
        raise ValueError("benchmark manifest is empty")
    return records


def select_completed_run(
    dataset: list[dict[str, Any]], manifest: list[dict[str, Any]]
) -> BenchmarkRun:
    """Select the latest run with all expected visual pairs and grammar skips."""
    by_run: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for entry in manifest:
        run_id = entry.get("run_id")
        if not isinstance(run_id, str) or not run_id:
            raise ValueError("manifest contains a record without a run_id")
        by_run[run_id].append(entry)

    completed: list[tuple[str, BenchmarkRun]] = []
    failures: list[str] = []
    for run_id, entries in by_run.items():
        try:
            run = _validate_run(run_id, dataset, entries)
        except ValueError as error:
            failures.append(f"{run_id}: {error}")
            continue
        latest_timestamp = max(_timestamp(entry) for entry in entries)
        completed.append((latest_timestamp, run))

    if not completed:
        details = "; ".join(failures) or "no run records found"
        raise ValueError(f"no complete benchmark run is available: {details}")
    return max(completed, key=lambda candidate: candidate[0])[1]


def _validate_run(
    run_id: str, dataset: list[dict[str, Any]], entries: list[dict[str, Any]]
) -> BenchmarkRun:
    visual_records = [record for record in dataset if record["visual_strategy"] != "grammar_function"]
    grammar_records = [record for record in dataset if record["visual_strategy"] == "grammar_function"]
    if len(visual_records) != 9 or len(grammar_records) != 1:
        raise ValueError("dataset must contain nine visual records and one grammar record")

    visual_entries: dict[str, dict[str, dict[str, Any]]] = {}
    used_paths: set[Path] = set()
    for record in visual_records:
        dataset_id = record["id"]
        visual_entries[dataset_id] = {}
        for model_id in MODEL_ORDER:
            matches = [
                entry
                for entry in entries
                if entry.get("dataset_id") == dataset_id
                and entry.get("model") == model_id
                and entry.get("status") == "SUCCEEDED"
            ]
            if len(matches) != 1:
                raise ValueError(
                    f"{dataset_id} requires exactly one SUCCEEDED record for {model_id}; "
                    f"found {len(matches)}"
                )
            image_path = Path(str(matches[0].get("output_path", "")))
            if not image_path.is_file():
                raise FileNotFoundError(f"expected image is missing: {image_path}")
            resolved_path = image_path.resolve()
            if resolved_path in used_paths:
                raise ValueError(f"duplicate generated image path in manifest: {image_path}")
            used_paths.add(resolved_path)
            visual_entries[dataset_id][model_id] = matches[0]

    grammar_id = grammar_records[0]["id"]
    grammar_entries = [entry for entry in entries if entry.get("dataset_id") == grammar_id]
    for model_id in MODEL_ORDER:
        matches = [
            entry
            for entry in grammar_entries
            if entry.get("model") == model_id and entry.get("status") == "SKIPPED_NON_VISUAL"
        ]
        if len(matches) != 1:
            raise ValueError(
                f"{grammar_id} requires exactly one SKIPPED_NON_VISUAL record for {model_id}; "
                f"found {len(matches)}"
            )

    if len(used_paths) != 18:
        raise ValueError(f"expected 18 unique generated images; found {len(used_paths)}")
    return BenchmarkRun(run_id, visual_entries, grammar_entries)


def _timestamp(entry: dict[str, Any]) -> str:
    timestamp = entry.get("timestamp")
    if not isinstance(timestamp, str):
        raise ValueError("manifest record has no timestamp")
    return timestamp


def generate_reports() -> BenchmarkRun:
    """Generate both human annotation deliverables from one validated benchmark run."""
    dataset = load_dataset()
    benchmark_run = select_completed_run(dataset, load_manifest())
    REPORTS_PATH.mkdir(parents=True, exist_ok=True)
    generate_docx(dataset, benchmark_run, DOCX_PATH)
    generate_pdf(dataset, benchmark_run, PDF_PATH)
    return benchmark_run


def generate_docx(dataset: list[dict[str, Any]], benchmark_run: BenchmarkRun, path: Path) -> None:
    """Create an A4 DOCX with one visual comparison per page."""
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.25)
    section.right_margin = Cm(1.25)
    _set_docx_defaults(document)
    _add_docx_page_number(section)

    title = document.add_heading("Ngaam-Nou Image Generation Benchmark", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("Human Annotation & Qwen vs Nano Banana 2 Comparison")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(14)
    document.add_paragraph("Step 6B - 18-image visual benchmark")
    document.add_paragraph(f"Benchmark run ID: {benchmark_run.run_id}")
    _add_docx_instructions(document)
    document.add_page_break()

    for record in _visual_records(dataset):
        _add_docx_visual_page(document, record, benchmark_run.visual_entries[record["id"]])
        document.add_page_break()

    _add_docx_grammar_page(document, _grammar_record(dataset), benchmark_run)
    document.add_page_break()
    _add_docx_summary(document, _visual_records(dataset))
    document.save(path)


def _set_docx_defaults(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")


def _add_docx_page_number(section: Any) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def _add_docx_instructions(document: Document) -> None:
    document.add_heading("Annotation Instructions", 1)
    document.add_paragraph("Score each criterion from 1-5. Do not calculate a winner or select one automatically.")
    document.add_paragraph("General criteria: 1 = very poor / unusable, 2 = weak, 3 = acceptable, 4 = good, 5 = excellent.")
    document.add_paragraph("Text contamination: 1 = severe unwanted text/writing contamination; 2 = noticeable unwanted text; 3 = minor contamination; 4 = almost clean; 5 = completely free of unwanted text.")
    document.add_paragraph("Images are paired only for the same vocabulary item. Complete notes and the final preference field as human judgments.")


def _add_docx_visual_page(
    document: Document, record: dict[str, Any], entries: dict[str, dict[str, Any]]
) -> None:
    document.add_heading(f"Vocabulary: {record['text']}", 1)
    document.add_paragraph(
        f"Meaning: {record['meaning']}\n"
        f"Category: {record['category']}\n"
        f"Visual strategy: {record['visual_strategy']}\n"
        f"Prompt version: {_shared_prompt_version(entries.values())}"
    )
    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    for column, model_id in enumerate(MODEL_ORDER):
        label = table.cell(0, column).paragraphs[0]
        label.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label.add_run(MODEL_LABELS[model_id]).bold = True
        image_paragraph = table.cell(1, column).paragraphs[0]
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.add_run().add_picture(str(entries[model_id]["output_path"]), width=Cm(7.7))
    document.add_paragraph()
    _add_docx_annotation_fields(document)


def _add_docx_annotation_fields(document: Document) -> None:
    document.add_heading("Human Annotations", 2)
    for number, criterion in enumerate(CRITERIA, 1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.add_run(f"{number}. {criterion}:  1   2   3   4   5\n").bold = True
        paragraph.add_run("Notes: " + "_" * 78)
    document.add_paragraph()
    overall = document.add_paragraph()
    overall.add_run("Overall annotation\n").bold = True
    overall.add_run("Qwen notes: " + "_" * 70 + "\n")
    overall.add_run("Nano Banana 2 notes: " + "_" * 61 + "\n")
    overall.add_run("Differences / observations: " + "_" * 56 + "\n")
    overall.add_run("Preferred for further testing (human annotation): " + "_" * 38)


def _add_docx_grammar_page(document: Document, record: dict[str, Any], benchmark_run: BenchmarkRun) -> None:
    document.add_heading(f"Vocabulary: {record['text']}", 1)
    document.add_paragraph(
        f"Category: {record['category']}\n"
        f"Visualizability: {record['visualizability']}\n"
        f"Strategy: {record['visual_strategy']}\n\n"
        "Status: SKIPPED_NON_VISUAL\n\n"
        "This benchmark intentionally does not force a literal image for a grammar/function word."
    )
    document.add_heading("Human Annotations", 2)
    document.add_paragraph("Alternative educational modality: " + "_" * 58)
    document.add_paragraph("Notes: " + "_" * 83)
    document.add_paragraph(f"Manifest records verified: {len(benchmark_run.grammar_entries)} skipped entries")


def _add_docx_summary(document: Document, visual_records: list[dict[str, Any]]) -> None:
    document.add_heading("Annotation Summary", 1)
    document.add_paragraph("Visual cases evaluated: 9\nGenerated images: 18\nModels: Qwen Image 2, Nano Banana 2\nNon-visual cases: 1 (的)")
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headers = ("Vocabulary", "Qwen avg.", "Nano avg.", "Qwen notes", "Nano notes", "Human decision")
    for cell, label in zip(table.rows[0].cells, headers):
        cell.text = label
    for record in visual_records:
        cells = table.add_row().cells
        cells[0].text = str(record["text"])
        for cell in cells[1:]:
            cell.text = ""
    document.add_paragraph("Leave all scores, notes, and decisions as human annotations. Do not calculate model rankings.")


def generate_pdf(dataset: list[dict[str, Any]], benchmark_run: BenchmarkRun, path: Path) -> None:
    """Create an A4 PDF with embedded source images and handwritten-note space."""
    _register_pdf_font()
    report = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    _draw_pdf_title_page(report, width, height, benchmark_run.run_id)
    for page_number, record in enumerate(_visual_records(dataset), 2):
        _draw_pdf_visual_page(
            report,
            width,
            height,
            page_number,
            record,
            benchmark_run.visual_entries[record["id"]],
        )
    _draw_pdf_grammar_page(report, width, height, 11, _grammar_record(dataset))
    _draw_pdf_summary_page(report, width, height, 12, _visual_records(dataset))
    report.save()


def _register_pdf_font() -> None:
    if not PDF_CJK_FONT_PATH.is_file():
        raise FileNotFoundError(f"Traditional Chinese PDF font not found: {PDF_CJK_FONT_PATH}")
    pdfmetrics.registerFont(TTFont(PDF_CJK_FONT, str(PDF_CJK_FONT_PATH)))


def _draw_pdf_title_page(report: canvas.Canvas, width: float, height: float, run_id: str) -> None:
    report.setFont("Helvetica-Bold", 20)
    report.drawCentredString(width / 2, height - 90, "Ngaam-Nou Image Generation Benchmark")
    report.setFont("Helvetica", 13)
    report.drawCentredString(width / 2, height - 116, "Human Annotation & Qwen vs Nano Banana 2 Comparison")
    report.drawCentredString(width / 2, height - 138, "Step 6B - 18-image visual benchmark")
    report.setFont("Helvetica", 10)
    report.drawString(48, height - 180, f"Benchmark run ID: {run_id}")
    report.setFont("Helvetica-Bold", 13)
    report.drawString(48, height - 218, "Annotation Instructions")
    report.setFont("Helvetica", 10)
    instructions = (
        "Score each criterion from 1-5. Do not calculate a winner or select one automatically.",
        "General criteria: 1 = very poor / unusable, 2 = weak, 3 = acceptable, 4 = good, 5 = excellent.",
        "Text contamination: 1 = severe unwanted text/writing contamination; 2 = noticeable unwanted text;",
        "3 = minor contamination; 4 = almost clean; 5 = completely free of unwanted text.",
        "Images are paired only for the same vocabulary item. Complete all notes as human judgments.",
    )
    _draw_pdf_lines(report, 48, height - 242, instructions, 15)
    _draw_pdf_page_number(report, width, 1)
    report.showPage()


def _draw_pdf_visual_page(
    report: canvas.Canvas,
    width: float,
    height: float,
    page_number: int,
    record: dict[str, Any],
    entries: dict[str, dict[str, Any]],
) -> None:
    report.setFont(PDF_CJK_FONT, 18)
    report.drawString(42, height - 48, f"Vocabulary: {record['text']}")
    report.setFont("Helvetica", 9)
    metadata = (
        f"Meaning: {record['meaning']}",
        f"Category: {record['category']}",
        f"Visual strategy: {record['visual_strategy']}",
        f"Prompt version: {_shared_prompt_version(entries.values())}",
    )
    _draw_pdf_lines(report, 42, height - 70, metadata, 12)
    image_size = 226
    left_x = 48
    right_x = width - 48 - image_size
    image_y = height - 365
    for x, model_id in ((left_x, MODEL_ORDER[0]), (right_x, MODEL_ORDER[1])):
        report.setFont("Helvetica-Bold", 11)
        report.drawCentredString(x + image_size / 2, image_y + image_size + 16, MODEL_LABELS[model_id])
        report.drawImage(
            ImageReader(str(entries[model_id]["output_path"])),
            x,
            image_y,
            width=image_size,
            height=image_size,
            preserveAspectRatio=True,
            anchor="c",
        )
        report.setStrokeColor(colors.HexColor("#8796A5"))
        report.rect(x, image_y, image_size, image_size, stroke=1, fill=0)
    report.setFont("Helvetica-Bold", 12)
    report.drawString(42, image_y - 22, "Human Annotations")
    y = image_y - 40
    report.setFont("Helvetica", 9)
    for number, criterion in enumerate(CRITERIA, 1):
        report.drawString(42, y, f"{number}. {criterion}:  1   2   3   4   5")
        report.drawString(270, y, "Notes:")
        report.line(305, y - 2, width - 42, y - 2)
        y -= 21
    report.setFont("Helvetica-Bold", 10)
    report.drawString(42, y - 2, "Overall annotation")
    report.setFont("Helvetica", 9)
    for label in (
        "Qwen notes:",
        "Nano Banana 2 notes:",
        "Differences / observations:",
        "Preferred for further testing (human annotation):",
    ):
        y -= 17
        report.drawString(42, y, label)
        report.line(230, y - 2, width - 42, y - 2)
    _draw_pdf_page_number(report, width, page_number)
    report.showPage()


def _draw_pdf_grammar_page(
    report: canvas.Canvas, width: float, height: float, page_number: int, record: dict[str, Any]
) -> None:
    report.setFont(PDF_CJK_FONT, 18)
    report.drawString(42, height - 48, f"Vocabulary: {record['text']}")
    report.setFont("Helvetica", 11)
    lines = (
        f"Category: {record['category']}",
        f"Visualizability: {record['visualizability']}",
        f"Strategy: {record['visual_strategy']}",
        "",
        "Status: SKIPPED_NON_VISUAL",
        "",
        "This benchmark intentionally does not force a literal image for a grammar/function word.",
    )
    _draw_pdf_lines(report, 42, height - 78, lines, 18)
    report.setFont("Helvetica-Bold", 12)
    report.drawString(42, height - 230, "Human Annotations")
    report.setFont("Helvetica", 10)
    report.drawString(42, height - 258, "Alternative educational modality:")
    report.line(205, height - 260, width - 42, height - 260)
    report.drawString(42, height - 292, "Notes:")
    for offset in (294, 314, 334):
        report.line(82, height - offset, width - 42, height - offset)
    _draw_pdf_page_number(report, width, page_number)
    report.showPage()


def _draw_pdf_summary_page(
    report: canvas.Canvas, width: float, height: float, page_number: int, visual_records: list[dict[str, Any]]
) -> None:
    report.setFont("Helvetica-Bold", 18)
    report.drawString(42, height - 48, "Annotation Summary")
    report.setFont(PDF_CJK_FONT, 10)
    _draw_pdf_lines(
        report,
        42,
        height - 74,
        ("Visual cases evaluated: 9", "Generated images: 18", "Models: Qwen Image 2, Nano Banana 2", "Non-visual cases: 1 (的)"),
        15,
    )
    headers = ("Vocabulary", "Qwen avg.", "Nano avg.", "Qwen notes", "Nano notes", "Human decision")
    widths = (68, 66, 66, 104, 104, 128)
    x = 42
    y = height - 165
    report.setFont("Helvetica-Bold", 8)
    for label, column_width in zip(headers, widths):
        report.rect(x, y, column_width, 22, stroke=1, fill=0)
        report.drawCentredString(x + column_width / 2, y + 8, label)
        x += column_width
    report.setFont(PDF_CJK_FONT, 10)
    for record in visual_records:
        x = 42
        y -= 38
        for index, column_width in enumerate(widths):
            report.rect(x, y, column_width, 38, stroke=1, fill=0)
            if index == 0:
                report.drawCentredString(x + column_width / 2, y + 14, str(record["text"]))
            x += column_width
    report.setFont("Helvetica", 9)
    report.drawString(42, y - 32, "Leave scores, notes, and decisions blank for human annotation. Do not calculate model rankings.")
    _draw_pdf_page_number(report, width, page_number)
    report.showPage()


def _draw_pdf_lines(
    report: canvas.Canvas, x: float, y: float, lines: Iterable[str], leading: float
) -> None:
    for line in lines:
        report.drawString(x, y, line)
        y -= leading


def _draw_pdf_page_number(report: canvas.Canvas, width: float, page_number: int) -> None:
    report.setFont("Helvetica", 8)
    report.drawCentredString(width / 2, 22, f"Page {page_number}")


def _visual_records(dataset: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [record for record in dataset if record["visual_strategy"] != "grammar_function"]


def _grammar_record(dataset: list[dict[str, Any]]) -> dict[str, Any]:
    grammar_records = [record for record in dataset if record["visual_strategy"] == "grammar_function"]
    if len(grammar_records) != 1:
        raise ValueError("expected exactly one grammar/function dataset record")
    return grammar_records[0]


def _shared_prompt_version(entries: Iterable[dict[str, Any]]) -> str:
    versions = {str(entry.get("prompt_version", "unknown")) for entry in entries}
    if len(versions) != 1:
        raise ValueError(f"paired image records use different prompt versions: {sorted(versions)}")
    return versions.pop()


def main() -> None:
    benchmark_run = generate_reports()
    print(f"run_id={benchmark_run.run_id}")
    print("visual_cases=9")
    print("embedded_images=18")
    print("non_visual_cases=1")
    print(f"docx={DOCX_PATH}")
    print(f"pdf={PDF_PATH}")


if __name__ == "__main__":
    main()